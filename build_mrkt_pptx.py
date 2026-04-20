from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

IMG_DIR = os.path.join(os.path.dirname(__file__), "images")
OUT = os.path.join(os.path.dirname(__file__), "MRKT_PLCE_Presentation.pptx")

AIR_FACTORY = [
    {
        "title": "Menthol",
        "subtitle": "Cool Menthol & Ice",
        "desc": "A crisp, clean menthol with a frosty icy finish. Classic and refreshing.",
        "image": "air-factory-menthol.jpg",
        "dual": True,
    },
    {
        "title": "Mint",
        "subtitle": "Fresh Mint & Ice",
        "desc": "Bright spearmint with a smooth cooling exhale. Pure and invigorating.",
        "image": "air-factory-mint.jpg",
        "dual": True,
    },
    {
        "title": "Unflavored",
        "subtitle": "No Ice — Pure Vapor",
        "desc": "Clean, smooth, and completely flavor-free. Just pure vapor with no additives.",
        "image": "air-factory-unflavored.jpg",
        "dual": False,
    },
]

SLIDES = [
    {
        "title": "Blood Orange Tangoberry",
        "subtitle": "Blood Orange, Grapefruit & Tangoberries",
        "desc": "A bright and vibrant mix of sweet blood orange, juicy grapefruit, and fresh tangoberries!",
        "image": "mrkt-plce-blood-orange-tangoberry.jpg",
    },
    {
        "title": "Forbidden Berry",
        "subtitle": "Tropical Fruit & Mixed Berries",
        "desc": "Perfectly captures the flavors of tropical fruit intermingled with a mixed berry combination – forbidden fruit tastes the best.",
        "image": "mrkt-plce-forbidden-berry.jpg",
    },
    {
        "title": "Fuji Pear Mangoberry",
        "subtitle": "Fuji Apple, Pear, Mango & Berries",
        "desc": "A blend of fresh fuji apples, sweet pears, ripe mangos, and succulent berries. A tropical fruit explosion!",
        "image": "mrkt-plce-fuji-pear-mangoberry.jpg",
    },
    {
        "title": "Pineapple Peach Dragonberry",
        "subtitle": "Pineapple, Peach & Dragonfruit",
        "desc": "Juicy pineapple infused with sweet peaches and a touch of dragonfruit.",
        "image": "mrkt-plce-pineapple-peach-dragonberry.jpg",
    },
    {
        "title": "Watermelon Hula Berry Lime",
        "subtitle": "Watermelon, Hula Berry & Lime",
        "desc": "An infusion of sweet watermelon, hula berry, and fresh-squeezed lime.",
        "image": "mrkt-plce-watermelon-hula-berry-lime.jpg",
    },
    {
        "title": "Thai Apple Melon Razz",
        "subtitle": "Apple, Melon & Raspberry",
        "desc": "Bold apple, lush melon, and zesty raspberry come together in a lively, flavorful vapor.",
        "image": "mrkt-plce-thai-apple-melon-razz.jpg",
    },
    {
        "title": "Iced Fuji Pear Mangoberry",
        "subtitle": "Fuji Apple, Pear, Mango & Berries — Iced",
        "desc": "All the tropical freshness of Fuji Pear Mangoberry, finished with a cool icy exhale.",
        "image": "mrkt-plce-iced-fuji-pear-mangoberry.jpg",
    },
    {
        "title": "Iced Forbidden Berry",
        "subtitle": "Tropical Fruit & Mixed Berries — Iced",
        "desc": "The beloved Forbidden Berry blend elevated with a refreshing burst of ice.",
        "image": "mrkt-plce-iced-forbidden-berry.jpg",
    },
]

# Widescreen 16:9
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]  # totally blank

# --- Cover slide ---
cover = prs.slides.add_slide(blank_layout)
bg = cover.background.fill
bg.solid()
bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

logo_path = os.path.join(IMG_DIR, "mrktplce.png")
if os.path.exists(logo_path):
    cover.shapes.add_picture(logo_path, Inches(4.16), Inches(1.5), Inches(5.0))

tf = cover.shapes.add_textbox(Inches(1), Inches(4.8), Inches(11.33), Inches(0.7))
p = tf.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "Available in Salt Nic (30ml) & Sub-Ohm (100ml)"
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
run.font.bold = False

tf2 = cover.shapes.add_textbox(Inches(1), Inches(5.6), Inches(11.33), Inches(0.55))
p2 = tf2.text_frame.paragraphs[0]
p2.alignment = PP_ALIGN.CENTER
run2 = p2.add_run()
run2.text = "The Vapor Place  ·  21+"
run2.font.size = Pt(14)
run2.font.color.rgb = RGBColor(0x88, 0x99, 0xAA)

# --- Product slides ---
for item in SLIDES:
    slide = prs.slides.add_slide(blank_layout)

    # White background
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    left_x = Inches(0.45)

    # Flavor title
    tf = slide.shapes.add_textbox(left_x, Inches(0.55), Inches(5.8), Inches(0.8))
    tf.text_frame.word_wrap = True
    p = tf.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = item["title"]
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    # Subtitle (gold)
    tf2 = slide.shapes.add_textbox(left_x, Inches(1.5), Inches(5.8), Inches(0.55))
    tf2.text_frame.word_wrap = True
    p2 = tf2.text_frame.paragraphs[0]
    run2 = p2.add_run()
    run2.text = item["subtitle"]
    run2.font.size = Pt(16)
    run2.font.color.rgb = RGBColor(0xCC, 0x99, 0x00)

    # Description
    tf3 = slide.shapes.add_textbox(left_x, Inches(2.2), Inches(5.8), Inches(1.5))
    tf3.text_frame.word_wrap = True
    p3 = tf3.text_frame.paragraphs[0]
    run3 = p3.add_run()
    run3.text = item["desc"]
    run3.font.size = Pt(15)
    run3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # Sizing badges area
    tf4 = slide.shapes.add_textbox(left_x, Inches(4.0), Inches(5.8), Inches(0.8))
    tf4.text_frame.word_wrap = True
    p4 = tf4.text_frame.paragraphs[0]
    run4a = p4.add_run()
    run4a.text = "Salt Nic  30ml  ·  24mg / 48mg     |     Sub-Ohm  100ml  ·  0mg / 3mg / 6mg"
    run4a.font.size = Pt(12)
    run4a.font.color.rgb = RGBColor(0x00, 0x88, 0xAA)

    # 21+ note
    tf5 = slide.shapes.add_textbox(left_x, Inches(5.0), Inches(5.8), Inches(0.5))
    p5 = tf5.text_frame.paragraphs[0]
    run5 = p5.add_run()
    run5.text = "21+  ·  The Vapor Place"
    run5.font.size = Pt(10)
    run5.font.color.rgb = RGBColor(0x88, 0x99, 0xAA)

    # Right panel: product image
    img_path = os.path.join(IMG_DIR, item["image"])
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(6.7), Inches(0.5), Inches(6.1), Inches(6.5))

# --- Air Factory divider slide ---
div = prs.slides.add_slide(blank_layout)
bg = div.background.fill
bg.solid()
bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
af_logo = os.path.join(IMG_DIR, "air-factory-menthol.jpg")  # placeholder — no brand logo file
tf = div.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11.33), Inches(1.2))
p = tf.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "AIR FACTORY"
run.font.size = Pt(54)
run.font.bold = True
run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
tf2 = div.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.33), Inches(0.6))
p2 = tf2.text_frame.paragraphs[0]
p2.alignment = PP_ALIGN.CENTER
run2 = p2.add_run()
run2.text = "Sub-Ohm 100ml  ·  Salt Nic 30ml"
run2.font.size = Pt(20)
run2.font.color.rgb = RGBColor(0x00, 0x88, 0xAA)

# --- Air Factory product slides ---
for item in AIR_FACTORY:
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    left_x = Inches(0.45)

    tf = slide.shapes.add_textbox(left_x, Inches(0.55), Inches(5.8), Inches(0.8))
    tf.text_frame.word_wrap = True
    p = tf.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Air Factory  " + item["title"]
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    tf2 = slide.shapes.add_textbox(left_x, Inches(1.5), Inches(5.8), Inches(0.55))
    tf2.text_frame.word_wrap = True
    p2 = tf2.text_frame.paragraphs[0]
    run2 = p2.add_run()
    run2.text = item["subtitle"]
    run2.font.size = Pt(16)
    run2.font.color.rgb = RGBColor(0xCC, 0x99, 0x00)

    tf3 = slide.shapes.add_textbox(left_x, Inches(2.2), Inches(5.8), Inches(1.5))
    tf3.text_frame.word_wrap = True
    p3 = tf3.text_frame.paragraphs[0]
    run3 = p3.add_run()
    run3.text = item["desc"]
    run3.font.size = Pt(15)
    run3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    tf4 = slide.shapes.add_textbox(left_x, Inches(4.0), Inches(5.8), Inches(0.8))
    tf4.text_frame.word_wrap = True
    p4 = tf4.text_frame.paragraphs[0]
    run4 = p4.add_run()
    if item["dual"]:
        run4.text = "Salt Nic  30ml  ·  24mg / 48mg     |     Sub-Ohm  100ml  ·  0mg / 3mg / 6mg"
    else:
        run4.text = "Sub-Ohm  100ml  ·  0mg / 3mg / 6mg"
    run4.font.size = Pt(12)
    run4.font.color.rgb = RGBColor(0x00, 0x88, 0xAA)

    tf5 = slide.shapes.add_textbox(left_x, Inches(5.0), Inches(5.8), Inches(0.5))
    p5 = tf5.text_frame.paragraphs[0]
    run5 = p5.add_run()
    run5.text = "21+  ·  The Vapor Place"
    run5.font.size = Pt(10)
    run5.font.color.rgb = RGBColor(0x88, 0x99, 0xAA)

    img_path = os.path.join(IMG_DIR, item["image"])
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(6.7), Inches(0.5), Inches(6.1), Inches(6.5))

prs.save(OUT)
print("Saved:", OUT)
